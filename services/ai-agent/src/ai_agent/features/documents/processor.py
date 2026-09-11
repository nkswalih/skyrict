"""Document content extraction for OCR pipeline (SKY-87).

Extracts text from document bytes fetched from core's storage backend.
Supports the same MIME types as the chat attachments extractor (PDF, DOCX,
XLSX, CSV, plain text) plus images (returned as-is for vision model).
"""

from __future__ import annotations

import base64
import csv
import io
from dataclasses import dataclass, field
from typing import Final

import structlog

logger = structlog.get_logger("ai_agent.documents")

_MAX_EXTRACT_BYTES: Final[int] = 10 * 1024 * 1024  # 10 MB

_IMAGE_MIMES: frozenset[str] = frozenset(
    {
        "image/jpeg",
        "image/png",
        "image/gif",
        "image/webp",
        "image/svg+xml",
    }
)

_TEXT_MIMES: frozenset[str] = frozenset(
    {
        "text/plain",
        "text/csv",
        "text/tab-separated-values",
        "text/markdown",
        "text/html",
        "application/json",
    }
)


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    """Result of extracting content from a document."""

    extracted_text: str = ""
    image_base64s: list[str] = field(default_factory=list)
    is_image: bool = False
    chunk_count: int = 0
    confidence_score: float | None = None


def _decode_text(raw: bytes) -> str:
    """Decode bytes to str using chardet detection with UTF-8 fallback."""
    try:
        import chardet

        result = chardet.detect(raw)
        encoding = result.get("encoding") or "utf-8"
    except ImportError:
        encoding = "utf-8"
    return raw.decode(encoding, errors="replace")


def extract_pdf(raw: bytes) -> str:
    """Extract text from a PDF file."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages: list[str] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Page {i + 1}]\n{text.strip()}")
    return "\n\n".join(pages) if pages else "(PDF contained no extractable text)"


def extract_docx(raw: bytes) -> str:
    """Extract text from a .docx file."""
    from docx import Document

    doc = Document(io.BytesIO(raw))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs) if paragraphs else "(Document contained no text)"


def extract_xlsx(raw: bytes) -> str:
    """Extract text from an .xlsx spreadsheet."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheets_text: list[str] = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append("\t".join(cells))
        if rows:
            sheets_text.append(f"[Sheet: {name}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(sheets_text) if sheets_text else "(Spreadsheet contained no data)"


def extract_csv_text(raw: bytes) -> str:
    """Extract text from a CSV/TSV file with auto-detected encoding."""
    text = _decode_text(raw)
    reader = csv.reader(io.StringIO(text))
    lines: list[str] = []
    for i, row in enumerate(reader):
        lines.append("\t".join(row))
        if i >= 500:
            lines.append(f"... ({i + 1}+ rows truncated)")
            break
    return "\n".join(lines) if lines else "(CSV contained no data)"


def extract_text_file(raw: bytes) -> str:
    """Extract text from a plain-text file."""
    return _decode_text(raw)


def extract_document_content(
    raw_bytes: bytes,
    mime_type: str,
    filename: str = "",
) -> ExtractionResult:
    """Extract text (and optionally images) from document bytes.

    Returns an ExtractionResult with extracted_text populated for text-bearing
    documents, or image_base64s for image files. chunk_count is estimated
    from text length (~400 tokens per chunk, ~4 chars per token).
    """
    if len(raw_bytes) > _MAX_EXTRACT_BYTES:
        logger.warning("document_too_large", filename=filename, size=len(raw_bytes))
        return ExtractionResult(
            extracted_text="(document too large for extraction)",
            chunk_count=0,
        )

    try:
        if mime_type in _IMAGE_MIMES:
            img_b64 = base64.b64encode(raw_bytes).decode("ascii")
            return ExtractionResult(
                extracted_text=f"[Image: {filename} - sent to vision model]",
                image_base64s=[img_b64],
                is_image=True,
                chunk_count=1,
            )
        elif mime_type == "application/pdf":
            text = extract_pdf(raw_bytes)
        elif mime_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ):
            text = extract_docx(raw_bytes)
        elif mime_type in ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",):
            text = extract_xlsx(raw_bytes)
        elif mime_type in ("text/csv", "text/tab-separated-values"):
            text = extract_csv_text(raw_bytes)
        elif mime_type in _TEXT_MIMES or mime_type.startswith("text/"):
            text = extract_text_file(raw_bytes)
        else:
            text = f"(unsupported file type: {mime_type})"

        # Estimate chunk count: ~4 chars/token, ~400 tokens/chunk
        char_count = len(text)
        chunk_count = max(1, char_count // 1600)

        return ExtractionResult(
            extracted_text=text,
            chunk_count=chunk_count,
        )
    except Exception:
        logger.exception("extraction_failed", filename=filename, mime=mime_type)
        return ExtractionResult(
            extracted_text=f"(extraction failed for {filename})",
            chunk_count=0,
        )


__all__ = [
    "ExtractionResult",
    "extract_csv_text",
    "extract_document_content",
    "extract_docx",
    "extract_pdf",
    "extract_text_file",
    "extract_xlsx",
]
